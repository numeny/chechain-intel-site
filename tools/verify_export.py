from openpyxl import load_workbook

for path in [
    "deliverables/australia_dealers_2026-10-23.xlsx",
    "deliverables/brazil_dealers_2026-10-23.xlsx",
    "deliverables/sample/australia_dealers_2026-10-23.xlsx",
    "deliverables/sample/brazil_dealers_2026-10-23.xlsx",
]:
    wb = load_workbook(path)
    print(path, wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(" ", name, ws.max_row - 1 if name != "说明" else ws.max_row)
