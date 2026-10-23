from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .data import BrandRelationRow, CompanyRow, GroupLinkRow, link_label, relationship_label, type_label

HEADERS = ["企业ID", "企业名称", "类型", "国家", "城市", "州/省", "地址", "代理品牌", "所属集团", "来源链接", "置信度", "更新时间"]
WIDTHS = [8, 32, 14, 8, 16, 10, 40, 24, 28, 45, 10, 12]
RELATION_HEADERS = ["企业ID", "企业名称", "品牌", "关系类型", "来源链接", "置信度", "更新时间"]
RELATION_WIDTHS = [8, 32, 16, 16, 45, 10, 12]
LINK_HEADERS = ["集团名称", "成员企业", "关系类型", "来源链接", "置信度"]
LINK_WIDTHS = [32, 32, 14, 45, 10]


def _style_sheet(ws, headers: list[str], widths: list[int], row_count: int) -> None:
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{max(row_count + 1, 1)}"


def _set_link(ws, row: int, column: int, url: str) -> None:
    cell = ws.cell(row=row, column=column)
    if url.startswith(("http://", "https://")):
        cell.hyperlink = url
        cell.style = "Hyperlink"
    cell.alignment = Alignment(wrap_text=True, vertical="top")


def write_workbook(
    rows: list[CompanyRow],
    path: str | Path,
    country_label: str,
    generated_at: str,
    relations: list[BrandRelationRow] | None = None,
    links: list[GroupLinkRow] | None = None,
    cta: str | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "企业"
    _style_sheet(ws, HEADERS, WIDTHS, len(rows))
    for r in rows:
        ws.append([
            r.company_id, r.name, type_label(r.entity_type), r.country.upper(),
            r.city, r.region, r.address, r.brands, r.group, r.source_url,
            r.confidence, r.updated_at,
        ])
        row = ws.max_row
        _set_link(ws, row, 10, r.source_url)
        ws.cell(row=row, column=7).alignment = Alignment(wrap_text=True, vertical="top")

    if relations is not None:
        rel_ws = wb.create_sheet("品牌代理明细")
        _style_sheet(rel_ws, RELATION_HEADERS, RELATION_WIDTHS, len(relations))
        for r in relations:
            rel_ws.append([
                r.company_id, r.company_name, r.brand, relationship_label(r.relationship_type),
                r.source_url, r.confidence, r.updated_at,
            ])
            _set_link(rel_ws, rel_ws.max_row, 5, r.source_url)

    if links is not None:
        link_ws = wb.create_sheet("集团关系")
        _style_sheet(link_ws, LINK_HEADERS, LINK_WIDTHS, len(links))
        for r in links:
            link_ws.append([
                r.parent_name, r.child_name, link_label(r.link_type),
                r.source_url, r.confidence,
            ])
            _set_link(link_ws, link_ws.max_row, 4, r.source_url)

    info = wb.create_sheet("说明")
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 80
    info_rows = [
        ("国家/地区", country_label),
        ("报告名称", f"{country_label}经销商与代理关系包"),
        ("生成日期", generated_at),
        ("数据口径", "仅包含已审核发布的公开来源记录；不含联系方式。"),
        ("字段说明", "代理品牌为该企业当前在售/代理品牌；所属集团来自公开集团关系。"),
        ("来源", "来源链接为公开网页或品牌方公开接口；抓取日期见各行更新时间。"),
        ("合规", "不含个人信息；如信息有误，可凭来源链接申请更正或删除。"),
        ("免责", "本报告为公开信息整理，不构成投资或交易建议；请以品牌方/企业官方信息为准。"),
    ]
    if cta:
        info_rows.append(("升级完整版", cta))
    for row in info_rows:
        info.append(row)
    for cell in info["A"]:
        cell.font = Font(bold=True)

    wb.save(path)
    return path
