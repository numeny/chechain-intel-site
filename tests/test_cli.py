import re
from openpyxl import load_workbook

from export_intel.cli import main


def test_cli_generates_full_and_sample(db_path, tmp_path):
    out = tmp_path / "deliverables"
    assert main(["--db", str(db_path), "--country", "au", "--kind", "full", "--out", str(out)]) == 0
    files = list(out.rglob("*"))
    assert any(p.suffix == ".xlsx" for p in files)
    assert any(p.suffix == ".pdf" for p in files)

    sample = out / "sample"
    assert main(["--db", str(db_path), "--country", "au", "--kind", "sample", "--out", str(sample)]) == 0
    xlsx = next(sample.rglob("*.xlsx"))
    wb = load_workbook(xlsx)
    ws = wb["企业"]
    assert ws.max_row - 1 <= 10
    addresses = [ws.cell(row=i, column=7).value or "" for i in range(2, ws.max_row + 1)]
    assert any(addresses)
    assert all(not re.search(r"\d", address) for address in addresses)
