from openpyxl import load_workbook

from export_intel.data import BrandRelationRow, CompanyRow, GroupLinkRow
from export_intel.excel import write_workbook


def _row(**overrides):
    data = dict(
        company_id=2, name="BYD Sydney", entity_type="dealer_store",
        country="au", city="Sydney", region="NSW",
        address="1 George St", brands="比亚迪", group="Autosports Group",
        source_url="https://example.com", confidence=0.9,
        updated_at="2026-10-17",
    )
    data.update(overrides)
    return CompanyRow(**data)


def test_write_workbook_creates_expected_sheets(tmp_path):
    path = write_workbook([_row()], tmp_path / "au.xlsx", "澳洲", "2026-10-23")
    assert path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == ["企业", "说明"]
    ws = wb["企业"]
    assert ws["A1"].value == "企业ID"
    assert ws["B2"].value == "BYD Sydney"
    assert ws["H2"].value == "比亚迪"
    assert ws["I2"].value == "Autosports Group"
    assert ws["J2"].hyperlink.target == "https://example.com"
    assert wb["说明"]["B1"].value == "澳洲"
    assert wb["说明"]["B2"].value == "澳洲经销商与代理关系包"


def test_write_workbook_adds_cta_row(tmp_path):
    path = write_workbook([_row()], tmp_path / "au.xlsx", "澳洲", "2026-10-23", cta="完整版 499 元")
    wb = load_workbook(path)
    values = [cell.value for cell in wb["说明"]["A"]]
    assert "升级完整版" in values


def test_write_workbook_adds_relation_and_link_sheets(tmp_path):
    relations = [BrandRelationRow(2, "BYD Sydney", "比亚迪", "official_dealer", "https://example.com", 0.9, "2026-10-17")]
    links = [GroupLinkRow("Autosports Group", "BYD Sydney", "group_member", "https://example.com", 0.7)]
    path = write_workbook(
        [_row()], tmp_path / "au.xlsx", "澳洲", "2026-10-23",
        relations=relations, links=links,
    )
    wb = load_workbook(path)
    assert wb.sheetnames == ["企业", "品牌代理明细", "集团关系", "说明"]
    assert wb["品牌代理明细"]["D2"].value == "官方授权经销商"
    assert wb["集团关系"]["C2"].value == "集团成员"
